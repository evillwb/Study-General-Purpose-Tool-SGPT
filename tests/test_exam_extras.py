"""Maths formulas, diagrams, exam formats and Kahoot/Quizlet exports."""
import io
import json
import time

import docx
import pymupdf
from openpyxl import load_workbook

from sgpt import convert, figures, quizapps
from sgpt.app import create_app
from sgpt.generator import EXAM_TEMPLATES, normalise_question


def _wait(c, job_id):
    for _ in range(300):
        job = c.get(f"/api/jobs/{job_id}").get_json()
        if job["state"] != "running":
            return job
        time.sleep(0.05)
    raise AssertionError("job did not finish")


def test_formulas_are_drawn_or_spelled_out(tmp_path):
    assert figures.math_image(r"\frac{3x+1}{2}")[:4] == b"\x89PNG"
    assert figures.math_image(r"\sqrt{b^2-4ac}", "svg").lstrip().startswith(b"<?xml")
    assert figures.math_image(r"\begin{matrix}1\end{matrix}") is None  # unsupported: no picture
    # A model that forgot to double its backslashes in JSON: "\f" and "\t" were eaten.
    assert figures.repair_math("Find $\x0crac{1}{2} \times 4$.") == r"Find $\frac{1}{2} \times 4$."
    assert figures.repair_math("Pay $5\nthen $3") == "Pay $5\nthen $3"
    assert figures.markdown_plain_math(r"Area $= \frac{1}{2} \times 6$ cm") == "Area = (1)/(2) × 6 cm"

    md = r"Solve $\frac{x}{2} = 3$." + "\n\nAnd $\\sqrt{16}$."
    with_images = figures.markdown_with_math_images(md, tmp_path)
    assert with_images.count("](m") == 2 and len(list(tmp_path.glob("m*.png"))) == 2
    data, _ = convert.export_markdown(with_images, "docx", "Maths", tmp_path)
    assert len(docx.Document(io.BytesIO(data)).inline_shapes) == 2  # formulas inside the lines
    data, _ = convert.export_markdown(with_images, "pdf", "Maths", tmp_path)
    page = pymupdf.open("pdf", data)[0]
    assert len(page.get_images()) == 2 and "Solve" in page.get_text()


def test_diagrams_are_checked_and_drawn():
    good = {"type": "shape", "points": {"A": [0, 0], "B": [4, 0], "C": [0, 3]},
            "sides": [["A", "B", "4 cm"]], "angles": {"A": "90°"}}
    assert figures.clean_diagram(good) == good
    assert figures.clean_diagram({"type": "graph", "functions": ["__import__('os').system('x')"]}) is None
    assert figures.clean_diagram({"type": "graph", "functions": ["open('x')"]}) is None
    assert figures.clean_diagram({"type": "pie"}) is None
    for spec in ({"type": "graph", "functions": ["x**2 - 4", "sin(x)"], "x_range": [-3, 3]},
                 {"type": "bar", "labels": ["A", "B"], "series": [{"values": [1, 2]}]},
                 {"type": "line", "labels": ["1", "2", "3"], "values": [3, 1, 2]},
                 {"type": "circle", "radius": 7}, {"type": "numberline", "range": [-2, 4]}):
        assert figures.diagram_image(spec)[:4] == b"\x89PNG", spec
    q = normalise_question({"question": "The diagram shows triangle ABC.", "type": "short", "diagram": good},
                           None, None)
    assert q["diagram"] == good


def test_exam_with_diagram_and_formula(tmp_path, fake_ai):
    host, _ = fake_ai
    c = create_app(tmp_path / "lib", tmp_path / "out", host).test_client()
    job = _wait(c, c.post("/api/generate", json={"topic": "Geometry", "structure": "custom",
                                                  "plan": [{"type": "short", "count": 2}]}).get_json()["id"])
    assert job["state"] == "done", job
    path = tmp_path / "out" / f"{job['exam_id']}.json"
    exam = json.loads(path.read_text())
    exam["questions"][0]["question"] = r"Find $\frac{3}{4}$ of the area."
    exam["questions"][0]["diagram"] = {"type": "shape", "points": {"A": [0, 0], "B": [4, 0], "C": [0, 3]}}
    path.write_text(json.dumps(exam))

    shown = c.get(f"/api/exams/{exam['id']}").get_json()
    assert '/api/math.svg?tex=%5Cfrac%7B3%7D%7B4%7D' in shown["html"]["paper"]
    assert f'/api/exams/{exam["id"]}/figure/diagram-1.svg' in shown["html"]["paper"]
    assert c.get(f"/api/exams/{exam['id']}/figure/diagram-1.svg").data.lstrip().startswith(b"<?xml")
    assert c.get(f"/api/exams/{exam['id']}/figure/diagram-2.svg").status_code == 404
    assert c.get("/api/math.svg?tex=x%5E2").status_code == 200

    word = docx.Document(io.BytesIO(c.get(f"/api/exams/{exam['id']}/export?part=paper&format=docx").data))
    assert len(word.inline_shapes) == 2  # the formula and the diagram
    md = c.get(f"/api/exams/{exam['id']}/export?part=paper&format=md").data.decode()
    assert "(3)/(4)" in md and "![Diagram" not in md


def test_exam_formats():
    assert {"spm_p1", "spm_p2", "uec_junior", "uec_senior", "igcse_p1", "igcse_p2", "stpm"} <= set(EXAM_TEMPLATES)
    assert sum(n for _, n, _ in EXAM_TEMPLATES["spm_p1"]["plan"]) == 40


def test_kahoot_and_quizlet(tmp_path, fake_ai):
    host, _ = fake_ai
    c = create_app(tmp_path / "lib", tmp_path / "out", host).test_client()
    job = _wait(c, c.post("/api/generate", json={"topic": "Light", "structure": "custom",
                                                  "plan": [{"type": "mcq", "count": 3}, {"type": "short", "count": 1}]}
                          ).get_json()["id"])
    r = c.get(f"/api/exams/{job['exam_id']}/export?format=kahoot")
    assert r.status_code == 200, r.data
    ws = load_workbook(io.BytesIO(r.data)).active
    assert ws["B8"].value.startswith("Question - max 120") and ws["H8"].value.startswith("Correct answer")
    rows = [[ws.cell(row=r_, column=col).value for col in range(1, 9)] for r_ in range(9, 20) if ws.cell(row=r_, column=2).value]
    assert len(rows) >= 2 and all(r_[6] == 30 and 1 <= r_[7] <= 4 for r_ in rows)
    text = c.get(f"/api/exams/{job['exam_id']}/export?format=quizlet").data.decode()
    assert len(text.strip().splitlines()) == 4 and all("\t" in line for line in text.strip().splitlines())

    data, cut = quizapps.kahoot_xlsx([{"question": "Q" * 200, "options": ["a", "b"], "answer": 1}], seconds=25)
    ws = load_workbook(io.BytesIO(data)).active
    assert len(ws["B9"].value) == 120 and cut == 1 and ws["G9"].value == 20 and ws["H9"].value == 2
    assert quizapps.quizlet_text([("osmosis", "movement of water\tacross a membrane")]) == \
        "osmosis\tmovement of water across a membrane\n"

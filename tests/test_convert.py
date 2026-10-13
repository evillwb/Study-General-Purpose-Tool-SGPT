import io
import time
import zipfile

import docx
import pymupdf
import pytest

from sgpt import convert
from sgpt.app import create_app

NOTES = "# Photosynthesis 光合作用\n\nPlants make **glucose**.\n\n- light\n- 水\n\n| A | B |\n|---|---|\n| 1 | 2 |\n"


def _scan(path, lines):
    page = pymupdf.open().new_page()
    for i, line in enumerate(lines):
        page.insert_text((50, 80 + 30 * i), line, fontsize=13)
    scan = pymupdf.open()
    scan.new_page().insert_image(scan[0].rect, stream=page.get_pixmap(dpi=150).tobytes("png"))
    scan.save(path)
    return path


def test_markdown_to_every_format(tmp_path):
    src = tmp_path / "notes.md"
    src.write_text(NOTES, encoding="utf-8")
    pdf = convert.convert(src, "pdf", tmp_path)
    assert "光合作用" in pymupdf.open(pdf)[0].get_text()
    word = docx.Document(convert.convert(src, "docx", tmp_path))
    assert word.paragraphs[0].text == "Photosynthesis 光合作用"
    assert any(r.bold and r.text == "glucose" for p in word.paragraphs for r in p.runs)
    assert word.tables[0].cell(1, 1).text == "2"
    assert "<strong>glucose</strong>" in convert.convert(src, "html", tmp_path).read_text()
    assert "**" not in convert.convert(src, "txt", tmp_path).read_text()


def test_word_round_trip_and_images(tmp_path):
    src = tmp_path / "notes.md"
    src.write_text(NOTES, encoding="utf-8")
    word = convert.convert(src, "docx", tmp_path)
    back = convert.docx_to_markdown(word)
    assert back.startswith("# Photosynthesis") and "**glucose**" in back and "| 1 | 2 |" in back
    (tmp_path / "out").mkdir()
    pdf = convert.convert(word, "pdf", tmp_path / "out")  # LibreOffice if present, else built in
    assert "Plants make" in pymupdf.open(pdf)[0].get_text()
    png = convert.convert(pdf, "png", tmp_path)
    assert png.suffix == ".png"
    two = convert.merge_pdfs([pdf, png], tmp_path / "merged.pdf", tmp_path / "out")
    assert len(pymupdf.open(two)) == 2
    assert zipfile.is_zipfile(convert.convert(two, "png", tmp_path))


def test_scan_to_word_and_searchable_pdf(tmp_path):
    scan = _scan(tmp_path / "scan.pdf", ["1. Define momentum. [1 mark]", "2. State Ohm's law. [2 marks]"])
    word = docx.Document(convert.convert(scan, "docx", tmp_path))  # OCR, not pdf2docx
    assert "momentum" in "\n".join(p.text for p in word.paragraphs)
    out = convert.searchable_pdf(scan, tmp_path / "searchable.pdf")
    page = pymupdf.open(out)[0]
    assert "Ohm" in page.get_text() and page.get_images()  # text layer on top of the scan


def test_unsupported_file(tmp_path):
    bad = tmp_path / "x.exe"
    bad.write_bytes(b"MZ")
    with pytest.raises(convert.ConvertError):
        convert.convert(bad, "pdf", tmp_path)


def test_convert_and_export_endpoints(tmp_path):
    c = create_app(tmp_path / "lib", tmp_path / "out").test_client()
    job = c.post("/api/convert", data={"target": "docx",
                                       "files": [(io.BytesIO(NOTES.encode()), "a.md")]}).get_json()
    deadline = time.time() + 30
    while job["state"] == "running" and time.time() < deadline:
        time.sleep(0.1)
        job = c.get(f"/api/jobs/{job['id']}").get_json()
    assert job["state"] == "done", job
    [result] = job["result"]
    r = c.get(result["url"])
    assert r.status_code == 200 and r.data[:2] == b"PK"
    assert c.get(f"/api/converted/{job['id']}/../../secret").status_code == 404
    assert c.post("/api/convert", data={"target": "pdf",
                                        "files": [(io.BytesIO(b"x"), "v.exe")]}).status_code == 400

    r = c.post("/api/export", json={"markdown": "# Hi\n\nText", "format": "pdf", "title": "OCR 结果"})
    assert r.status_code == 200 and r.data[:4] == b"%PDF"
    assert "filename*=UTF-8''OCR%20%E7%BB%93%E6%9E%9C.pdf" in r.headers["Content-Disposition"]


def _convert(c, name, data, target="docx"):
    job = c.post("/api/convert", data={"target": target, "files": [(io.BytesIO(data), name)]}).get_json()
    deadline = time.time() + 30
    while job["state"] == "running" and time.time() < deadline:
        time.sleep(0.1)
        job = c.get(f"/api/jobs/{job['id']}").get_json()
    assert job["state"] == "done", job
    return job["result"][0]


def test_converted_files_can_be_listed_and_deleted(tmp_path):
    c = create_app(tmp_path / "lib", tmp_path / "out").test_client()
    first = _convert(c, "a.md", NOTES.encode())
    second = _convert(c, "b.md", NOTES.encode())
    listed = c.get("/api/converted").get_json()
    assert [f["name"] for f in listed] == ["b.docx", "a.docx"]  # newest first
    assert not list((tmp_path / "out" / "converted").glob("*/in"))  # uploads are removed

    assert c.delete(first["url"]).get_json()["ok"]
    assert [f["name"] for f in c.get("/api/converted").get_json()] == ["b.docx"]
    assert c.get(first["url"]).status_code == 404
    _convert(c, "c.md", NOTES.encode())
    assert c.delete("/api/converted").get_json()["removed"] == 2
    assert c.get("/api/converted").get_json() == [] and c.get(second["url"]).status_code == 404


def test_delete_all_data(tmp_path, monkeypatch):
    from sgpt.providers import config_path, save_provider_config

    lib, out = tmp_path / "lib", tmp_path / "out"
    c = create_app(lib, out).test_client()
    _convert(c, "a.md", NOTES.encode())
    (lib / "paper.pdf").write_bytes(b"%PDF-1.4")
    (lib / ".question_bank.json").write_text("{}")
    (out / "exam1.json").write_text("{}")
    save_provider_config("openai", api_key="sk-test")
    assert c.post("/api/data/delete-all", json={"confirm": "yes"}).status_code == 400
    assert (lib / "paper.pdf").exists()

    assert c.post("/api/data/delete-all", json={"confirm": "DELETE"}).get_json()["ok"]
    assert list(lib.iterdir()) == [] and [p.name for p in out.iterdir()] == ["study"]
    assert config_path().exists()  # settings are kept unless asked
    assert c.post("/api/data/delete-all", json={"confirm": "DELETE", "settings": True}).get_json()["ok"]
    assert not config_path().exists()
    assert c.get("/api/exams").status_code == 200


def test_running_jobs_report_time_left(tmp_path):
    from sgpt.app import Job

    job = Job("convert")
    job.started -= 20
    assert job.eta() is None  # nothing finished yet
    job.step(1, 5, "one of five")
    assert job.eta() == 80  # 20 s per step, 4 steps left
    job.state = "done"
    assert job.eta() is None
    c = create_app(tmp_path / "lib", tmp_path / "out").test_client()
    assert c.get("/api/jobs/running").get_json() == []


def _scan_with_diagram(path):
    page = pymupdf.open().new_page()
    page.insert_text((50, 70), "Bab 2 Geometri", fontsize=16)
    page.insert_text((50, 110), "1. The diagram shows triangle ABC. [4 markah]", fontsize=12)
    page.draw_polyline([(100, 140), (300, 140), (200, 260), (100, 140)], color=(0, 0, 0), width=2)
    page.insert_text((88, 150), "A", fontsize=12)
    page.insert_text((305, 150), "B", fontsize=12)
    page.insert_text((195, 280), "C", fontsize=12)
    page.insert_text((50, 330), "2. Find the area of the triangle. [2 markah]", fontsize=12)
    scan = pymupdf.open()
    scan.new_page().insert_image(scan[0].rect, stream=page.get_pixmap(dpi=150).tobytes("png"))
    scan.save(path)
    return path


def test_scan_to_word_keeps_the_layout(tmp_path):
    from sgpt.layout import layout_pdf

    scan = _scan_with_diagram(tmp_path / "scan.pdf")
    rebuilt = pymupdf.open(layout_pdf(scan, tmp_path / "rebuilt.pdf"))[0]
    words = {w[4]: w for w in rebuilt.get_text("words")}
    assert words["Geometri"][1] < words["triangle"][1] < words["area"][1]  # same order on the page
    assert abs(words["Bab"][1] - 55) < 12  # and about the same place (y in points)
    assert len(rebuilt.get_images()) == 1  # the diagram, as a picture
    assert "A" not in words and "C" not in words  # its labels stay in the picture

    notes, steps = [], []
    word = docx.Document(convert.convert(scan, "docx", tmp_path, notes=notes,
                                         step=lambda n, total: steps.append(n)))
    text = "\n".join(p.text for p in word.paragraphs)
    assert "triangle ABC" in text and "Find the area" in text
    assert len(word.inline_shapes) == 1 and steps == [1]


def test_word_to_pdf_without_office_keeps_tables_and_pictures(tmp_path, monkeypatch):
    monkeypatch.setattr(convert, "libreoffice", lambda: None)
    monkeypatch.setattr(convert, "ms_office", lambda: set())
    picture = tmp_path / "pic.png"
    pymupdf.open().new_page(width=120, height=80).get_pixmap().save(picture)
    d = docx.Document()
    d.add_heading("Cell structure", 1)
    d.add_paragraph("The nucleus controls the cell.")
    table = d.add_table(rows=2, cols=2)
    table.cell(0, 0).text, table.cell(1, 1).text = "Organelle", "Mitochondria"
    d.add_picture(str(picture))
    d.save(tmp_path / "notes.docx")
    notes = []
    pdf = convert.convert(tmp_path / "notes.docx", "pdf", tmp_path, notes=notes)
    page = pymupdf.open(pdf)[0]
    assert "Cell structure" in page.get_text() and "Mitochondria" in page.get_text()
    assert page.get_images()
    assert notes and "LibreOffice" in notes[0]
    assert convert.targets_for(".pptx") == []  # needs Office or LibreOffice

"""Export to Kahoot! (spreadsheet import) and Quizlet (copy-and-paste import).

Kahoot: Create → Add question → Import spreadsheet, using Kahoot's template layout (the
question in column B from row 9, four answers in C–F, the time limit in G and the number of
the correct answer in H). Only multiple-choice questions with 2–4 options fit.

Quizlet: Create → Import → paste the text. Each line is "term<TAB>definition".
"""

from __future__ import annotations

import io

from .figures import markdown_plain_math

KAHOOT_QUESTION, KAHOOT_ANSWER = 120, 75  # Kahoot's character limits
KAHOOT_TIMES = (5, 10, 20, 30, 60, 90, 120, 240)


def _short(text: str, limit: int) -> tuple[str, bool]:
    text = " ".join(markdown_plain_math(str(text)).replace("**", "").split())
    return (text, False) if len(text) <= limit else (text[:limit - 1] + "…", True)


def mcq_from_exam(exam: dict) -> list[dict]:
    """Multiple-choice questions of an exam: {"question", "options", "answer": index}."""
    out = []
    for q in exam["questions"]:
        if q["type"] not in ("mcq", "true_false") or len(q["options"]) < 2:
            continue
        key = (q.get("final_answer") or q.get("answer") or "").strip()
        index = next((i for i, o in enumerate(q["options"])
                      if key[:1].upper() == chr(65 + i) and (len(key) == 1 or not key[1:2].isalnum())
                      or (o and o.lower() in key.lower())), None)
        if index is not None:
            out.append({"question": q["question"], "options": q["options"], "answer": index})
    return out


def kahoot_xlsx(items: list[dict], seconds: int = 30) -> tuple[bytes, int]:
    """Kahoot import spreadsheet for MCQ items. Returns (xlsx bytes, how many texts were cut)."""
    from openpyxl import Workbook
    from openpyxl.styles import Font

    seconds = min(KAHOOT_TIMES, key=lambda t: abs(t - seconds))
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws["B2"] = "Quiz template"
    ws["B2"].font = Font(bold=True, size=14)
    ws["B4"] = "Made with SGPT. Import it in Kahoot!: Add question → Import spreadsheet."
    headers = ["Question - max 120 characters", "Answer 1 - max 75 characters",
               "Answer 2 - max 75 characters", "Answer 3 - max 75 characters",
               "Answer 4 - max 75 characters",
               "Time limit (sec) – 5, 10, 20, 30, 60, 90, 120, or 240 secs",
               "Correct answer(s) - choose at least one"]
    for col, header in enumerate(headers, start=2):
        ws.cell(row=8, column=col, value=header).font = Font(bold=True)
    cut = 0
    row = 9
    for item in items:
        options = item["options"][:4]
        if len(options) < 2 or not 0 <= item["answer"] < len(options):
            continue
        question, was_cut = _short(item["question"], KAHOOT_QUESTION)
        cut += was_cut
        ws.cell(row=row, column=1, value=row - 8)
        ws.cell(row=row, column=2, value=question)
        for i, option in enumerate(options):
            text, was_cut = _short(option, KAHOOT_ANSWER)
            cut += was_cut
            ws.cell(row=row, column=3 + i, value=text)
        ws.cell(row=row, column=7, value=seconds)
        ws.cell(row=row, column=8, value=item["answer"] + 1)
        row += 1
    if row == 9:
        raise ValueError("There are no multiple-choice questions to send to Kahoot.")
    for col, width in zip("BCDEFGH", (60, 28, 28, 28, 28, 18, 18)):
        ws.column_dimensions[col].width = width
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue(), cut


def quizlet_text(pairs: list[tuple[str, str]]) -> str:
    """Quizlet import text: one "term<TAB>definition" per line (tabs and line breaks removed)."""
    def clean(text):
        return " ".join(markdown_plain_math(str(text)).replace("**", "").replace("\t", " ").split())
    lines = [f"{clean(a)}\t{clean(b)}" for a, b in pairs if clean(a) and clean(b)]
    if not lines:
        raise ValueError("There is nothing to send to Quizlet.")
    return "\n".join(lines) + "\n"


def pairs_from_exam(exam: dict) -> list[tuple[str, str]]:
    out = []
    for q in exam["questions"]:
        question = q["question"]
        if q["options"]:
            question += " " + " ".join(f"{chr(65 + i)}. {o}" for i, o in enumerate(q["options"]))
        out.append((question, q.get("final_answer") or q.get("answer") or ""))
    return out


def pairs_from_study(s: dict) -> list[tuple[str, str]]:
    if s["kind"] == "flashcards":
        return [(c["front"], c["back"]) for c in s["cards"]]
    if s["kind"] == "vocab":
        return [(w["word"], w["meaning"] + (f" ({w['translation']})" if w.get("translation") else ""))
                for w in s["words"]]
    out = []
    for it in s.get("items") or []:
        if it["type"] == "mcq":
            out.append((it["question"] + " " + " / ".join(it["options"]), it["options"][it["answer"]]))
        elif it["type"] == "fill":
            out.append((it["sentence"] + (f" ({it['base']})" if it.get("base") else ""), it["answer"]))
        elif it["type"] == "error":
            out.append((" | ".join(it["segments"]), f"{chr(65 + it['answer'])} → {it['correction']}"))
    return out


def mcq_from_study(s: dict) -> list[dict]:
    return [{"question": it["question"], "options": it["options"], "answer": it["answer"]}
            for it in s.get("items") or [] if it.get("type") == "mcq"]

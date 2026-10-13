"""Convert documents between common formats, and export Markdown/text to Word, PDF or HTML.

Conversions keep the original look where they can:

- Word, PowerPoint and Excel -> PDF use Microsoft Office when it is installed (Windows),
  otherwise LibreOffice (free: https://www.libreoffice.org). Without either, Word files
  are laid out from their content (headings, tables, pictures), which is close but not exact.
- PDF -> Word keeps the page layout (pdf2docx).
- Scanned PDFs and photos -> Word are rebuilt page by page: text where it was, at its
  size, and diagrams cut out of the scan (see layout.py).
"""

from __future__ import annotations

import html as html_lib
import re
import shutil
import subprocess
import sys
import tempfile
import zipfile
from pathlib import Path

import markdown as md_lib

PDF = {".pdf"}
IMAGES = {".png", ".jpg", ".jpeg", ".webp", ".bmp", ".tif", ".tiff", ".gif"}
WORD = {".docx"}
OFFICE = {".doc", ".rtf", ".odt", ".pptx", ".ppt", ".odp", ".xlsx", ".xls", ".ods"}
MARKDOWN = {".md", ".markdown"}
TEXT = {".txt"}
HTML = {".html", ".htm"}
INPUTS = PDF | IMAGES | WORD | OFFICE | MARKDOWN | TEXT | HTML
TARGETS = {"pdf": "PDF", "docx": "Word (.docx)", "txt": "Text (.txt)", "md": "Markdown (.md)",
           "html": "Web page (.html)", "png": "Images (.png)"}

PAGE_CSS = """
body { font-family: serif; font-size: 10.5pt; line-height: 1.45; color: #111; }
h1 { font-size: 22pt; margin: 0 0 8pt; color: #1f3f7a; }
h2 { font-size: 15pt; margin: 16pt 0 6pt; color: #1f3f7a; border-bottom: 0.8pt solid #1f3f7a; padding-bottom: 2pt; }
h3 { font-size: 12pt; margin: 12pt 0 4pt; } h4 { font-size: 11pt; margin: 10pt 0 3pt; }
p { margin: 0 0 6pt; }
ul, ol { margin: 0 0 6pt 14pt; } li { margin: 0 0 2pt; }
blockquote { margin: 6pt 0 8pt; padding: 6pt 9pt; background-color: #f2f3f5; border: 0.6pt solid #b9bec7; }
blockquote p { margin: 0 0 2pt; }
table { border-collapse: collapse; margin: 6pt 0 8pt; width: 100%; }
th { border-top: 1pt solid #333; border-bottom: 0.6pt solid #333; padding: 3pt 5pt; text-align: left; }
td { border-bottom: 0.4pt solid #ccc; padding: 3pt 5pt; vertical-align: top; }
code, pre { font-family: monospace; font-size: 9.5pt; } hr { border: 0; border-top: 0.5pt solid #aaa; }
"""


class ConvertError(RuntimeError):
    pass


def targets_for(suffix: str) -> list[str]:
    suffix = suffix.lower()
    if suffix not in INPUTS:
        return []
    if suffix in OFFICE and not libreoffice():
        # Old Word files, PowerPoint and Excel need Microsoft Office or LibreOffice.
        app = _OFFICE_APP.get(suffix)
        if not (app and app in ms_office()):
            return []
        return ["pdf", "docx"] if app == "Word" else ["pdf", "png"]
    return [t for t in TARGETS if not (suffix in PDF and t == "pdf")]


# --- Microsoft Office (Windows): exact conversions with the real program ----------------
_OFFICE_APP = {".doc": "Word", ".docx": "Word", ".rtf": "Word", ".odt": "Word",
               ".ppt": "PowerPoint", ".pptx": "PowerPoint", ".odp": "PowerPoint",
               ".xls": "Excel", ".xlsx": "Excel", ".ods": "Excel"}


def ms_office() -> set[str]:
    """Which Microsoft Office programs can be driven on this computer (Windows only)."""
    if sys.platform != "win32":
        return set()
    try:
        import win32com.client  # noqa: F401  (pywin32)
        import winreg
    except ImportError:
        return set()
    found = set()
    for app in ("Word", "PowerPoint", "Excel"):
        try:
            winreg.CloseKey(winreg.OpenKey(winreg.HKEY_CLASSES_ROOT, f"{app}.Application"))
            found.add(app)
        except OSError:
            pass
    return found


def _ms_office_convert(src: Path, fmt: str, out: Path) -> Path | None:
    """Convert with Microsoft Office itself. None if Office isn't there or it fails."""
    app_name = _OFFICE_APP.get(src.suffix.lower())
    if not app_name or app_name not in ms_office() or (fmt != "pdf" and app_name != "Word"):
        return None
    import pythoncom
    import win32com.client

    source, target = str(src.resolve()), str(out.resolve())
    pythoncom.CoInitialize()  # this runs in a background thread
    app = None
    try:
        app = win32com.client.DispatchEx(f"{app_name}.Application")
        if app_name == "Word":
            app.Visible, app.DisplayAlerts = False, 0
            doc = app.Documents.Open(source, ReadOnly=True, AddToRecentFiles=False, Visible=False)
            try:
                doc.SaveAs2(target, FileFormat=17 if fmt == "pdf" else 16)  # 17 PDF, 16 .docx
            finally:
                doc.Close(False)
        elif app_name == "PowerPoint":
            pres = app.Presentations.Open(source, ReadOnly=True, Untitled=False, WithWindow=False)
            try:
                pres.SaveAs(target, 32)  # 32 = PDF
            finally:
                pres.Close()
        else:
            app.Visible, app.DisplayAlerts = False, False
            book = app.Workbooks.Open(source, ReadOnly=True)
            try:
                book.ExportAsFixedFormat(0, target)  # 0 = PDF
            finally:
                book.Close(False)
    except Exception:  # noqa: BLE001 - any Office automation error: try the next way
        return None
    finally:
        if app is not None:
            try:
                app.Quit()
            except Exception:  # noqa: BLE001
                pass
        pythoncom.CoUninitialize()
    return out if out.exists() else None


def _office_convert(src: Path, fmt: str, out: Path) -> Path | None:
    """Microsoft Office, then LibreOffice. None when neither is installed."""
    done = _ms_office_convert(src, fmt, out)
    if done:
        return done
    if libreoffice():
        converted = _lo_convert(src, fmt, out.parent)
        return converted if converted == out else Path(shutil.move(converted, out))
    return None


# --- Building blocks --------------------------------------------------------------------

def libreoffice() -> str | None:
    """Path to LibreOffice's soffice, if installed."""
    for name in ("soffice", "libreoffice"):
        if shutil.which(name):
            return shutil.which(name)
    for path in (r"C:\Program Files\LibreOffice\program\soffice.exe",
                 r"C:\Program Files (x86)\LibreOffice\program\soffice.exe",
                 "/Applications/LibreOffice.app/Contents/MacOS/soffice"):
        if Path(path).exists():
            return path
    return None


def _lo_convert(src: Path, fmt: str, out_dir: Path) -> Path:
    soffice = libreoffice()
    if not soffice:
        raise ConvertError(f"Converting {src.suffix} files needs LibreOffice (free): "
                           "https://www.libreoffice.org/download/")
    with tempfile.TemporaryDirectory() as profile:
        # A private profile lets this run even while LibreOffice is open.
        cmd = [soffice, f"-env:UserInstallation=file:///{Path(profile).as_posix().lstrip('/')}",
               "--headless", "--convert-to", fmt, "--outdir", str(out_dir), str(src)]
        try:
            subprocess.run(cmd, check=True, capture_output=True, timeout=240)
        except (subprocess.SubprocessError, OSError) as exc:
            raise ConvertError(f"LibreOffice could not convert {src.name}: {exc}") from exc
    out = out_dir / f"{src.stem}.{fmt.split(':')[0]}"
    if not out.exists():
        raise ConvertError(f"LibreOffice could not convert {src.name}")
    return out


def keep_box_lines(text: str) -> str:
    """Each line of a "> " box stays on its own line (Markdown would join them)."""
    lines = text.split("\n")
    for i in range(len(lines) - 1):
        if lines[i].startswith(">") and lines[i + 1].startswith(">") and lines[i + 1][1:].strip() \
                and lines[i][1:].strip() and not lines[i].endswith("  "):
            lines[i] += "  "
    return "\n".join(lines)


def markdown_to_html(text: str, title: str = "") -> str:
    body = md_lib.markdown(keep_box_lines(text), extensions=["extra", "sane_lists"])
    return (f"<!doctype html><html><head><meta charset='utf-8'><title>{html_lib.escape(title)}"
            f"</title><style>{PAGE_CSS}</style></head><body>{body}</body></html>")


_IMG = re.compile(r'<img\b[^>]*\bsrc="([^"]+)"[^>]*>', re.IGNORECASE)
_MD_IMG = re.compile(r"^!\[([^\]]*)\]\(([^)\s]+)\)$")
MAX_IMAGE_WIDTH = 480  # points: the A4 text width used below


def _sized_images(html: str, images_dir: Path) -> str:
    """Give each <img> its print size (pictures are rendered at 2x), at most the text width."""
    import pymupdf

    def size(match):
        try:
            with pymupdf.open(images_dir / match.group(1)) as img:
                width = img[0].rect.width
        except (RuntimeError, OSError, IndexError):
            return match.group(0)
        return f'<img src="{match.group(1)}" width="{min(width / 2, MAX_IMAGE_WIDTH):.0f}">'
    return _IMG.sub(size, html)


def html_to_pdf(page_html: str, out: Path, header: str = "", images_dir: Path | None = None) -> Path:
    """Lay out HTML on A4 pages; with a header, every page after the first gets a running
    header and all pages a centred page number, like a printed handout. Images are read
    from images_dir."""
    import pymupdf

    body = re.search(r"<body[^>]*>(.*)</body>", page_html, re.DOTALL | re.IGNORECASE)
    html = body.group(1) if body else page_html
    if images_dir:
        html = _sized_images(html, images_dir)
    story = pymupdf.Story(html=html, user_css=PAGE_CSS,
                          archive=pymupdf.Archive(str(images_dir)) if images_dir else None)
    writer = pymupdf.DocumentWriter(str(out))
    page = pymupdf.paper_rect("a4")
    more = True
    while more:
        device = writer.begin_page(page)
        more, _ = story.place(page + (56, 62, -56, -56))
        story.draw(device)
        writer.end_page()
    writer.close()
    if header:
        doc = pymupdf.open(out)
        if len(doc) > 1:
            for number, pg in enumerate(doc, 1):
                width = pg.rect.width
                if number > 1:
                    pg.insert_htmlbox(pymupdf.Rect(56, 26, width - 56, 44),
                                      f"<div style='font-size:8.5pt;color:#555;font-family:serif'>"
                                      f"{html_lib.escape(header)}</div>")
                    pg.draw_line((56, 44), (width - 56, 44), color=(0.6, 0.6, 0.6), width=0.4)
                pg.insert_htmlbox(pymupdf.Rect(56, pg.rect.height - 38, width - 56, pg.rect.height - 24),
                                  f"<div style='font-size:9pt;text-align:center;font-family:serif'>"
                                  f"– {number} –</div>")
            doc.saveIncr()
        doc.close()
    return out


_INLINE = re.compile(r"(!\[[^\]]*\]\([^)\s]+\)|\*\*[^*]+\*\*|\*[^*\s][^*]*\*|`[^`]+`)")
_INLINE_IMG = re.compile(r"^!\[([^\]]*)\]\(([^)\s]+)\)$")


def _add_runs(paragraph, text: str, images_dir: Path | None = None) -> None:
    for piece in _INLINE.split(text):
        if not piece:
            continue
        picture = _INLINE_IMG.match(piece)
        if picture:  # a formula drawn as a picture, inside the line
            path = images_dir / picture.group(2) if images_dir else None
            if path and path.is_file():
                import pymupdf
                from docx.shared import Pt

                with pymupdf.open(path) as img:
                    height = img[0].rect.height / 2
                paragraph.add_run().add_picture(str(path), height=Pt(min(height, 60)))
            else:
                paragraph.add_run(picture.group(1))
        elif piece.startswith("**") and piece.endswith("**") and len(piece) > 4:
            paragraph.add_run(piece[2:-2]).bold = True
        elif piece.startswith("`") and piece.endswith("`"):
            paragraph.add_run(piece[1:-1]).font.name = "Consolas"
        elif piece.startswith("*") and piece.endswith("*") and len(piece) > 2:
            paragraph.add_run(piece[1:-1]).italic = True
        else:
            paragraph.add_run(piece)


def markdown_to_docx(text: str, out: Path, title: str = "", images_dir: Path | None = None) -> Path:
    """Write Markdown (headings, lists, quotes, tables, bold/italic, pictures on their own
    line) as a Word document."""
    import docx
    from docx.shared import Pt

    doc = docx.Document()
    if title:
        doc.core_properties.title = title
    lines = text.replace("\r\n", "\n").split("\n")
    i = 0
    while i < len(lines):
        line = lines[i].rstrip()
        stripped = line.strip()
        if not stripped or stripped in ("---", "***"):
            i += 1
            continue
        if stripped.startswith("|") and i + 1 < len(lines) and re.match(r"^\s*\|?\s*:?-{3}", lines[i + 1]):
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                if not re.match(r"^\s*\|?\s*:?-{3}", lines[i]):
                    rows.append([c.strip() for c in lines[i].strip().strip("|").split("|")])
                i += 1
            table = doc.add_table(rows=len(rows), cols=max(len(r) for r in rows))
            table.style = "Table Grid"
            for r, row in enumerate(rows):
                for c, cell in enumerate(row):
                    _add_runs(table.cell(r, c).paragraphs[0], cell, images_dir)
            continue
        picture = _MD_IMG.match(stripped)
        if picture and images_dir and (images_dir / picture.group(2)).is_file():
            import pymupdf

            path = images_dir / picture.group(2)
            with pymupdf.open(path) as img:
                width = img[0].rect.width / 2
            doc.add_picture(str(path), width=Pt(min(width, 450)))
            i += 1
            continue
        heading = re.match(r"^(#{1,6})\s+(.*)", stripped)
        if heading:
            level = min(len(heading.group(1)), 4)
            doc.add_heading(re.sub(r"\*\*(.+?)\*\*", r"\1", heading.group(2)), level=level - 1 if level > 1 else 0)
        elif re.match(r"^[-*+]\s+", stripped):
            _add_runs(doc.add_paragraph(style="List Bullet"), re.sub(r"^[-*+]\s+", "", stripped), images_dir)
        elif re.match(r"^\d+[.)]\s+", stripped):
            _add_runs(doc.add_paragraph(style="List Number"), re.sub(r"^\d+[.)]\s+", "", stripped), images_dir)
        elif stripped.startswith(">"):
            p = doc.add_paragraph(style="Quote")
            _add_runs(p, stripped.lstrip("> "), images_dir)
        else:
            # Join hard-wrapped lines ("  " at the end) into one paragraph with line breaks.
            p = doc.add_paragraph()
            parts = [stripped]
            while line.endswith("  ") and i + 1 < len(lines) and lines[i + 1].strip():
                i += 1
                line = lines[i].rstrip("\n")
                parts.append(lines[i].strip())
            for n, part in enumerate(parts):
                if n:
                    p.add_run().add_break()
                _add_runs(p, part, images_dir)
        i += 1
    doc.save(out)
    return out


def docx_to_markdown(path: Path) -> str:
    import docx

    doc = docx.Document(path)
    out = []
    for block in doc.element.body.iterchildren():
        tag = block.tag.rsplit("}", 1)[-1]
        if tag == "p":
            para = docx.text.paragraph.Paragraph(block, doc)
            text = "".join(f"**{r.text}**" if r.bold and r.text.strip() else r.text for r in para.runs)
            text = text.strip()
            if not text:
                continue
            style = (para.style.name or "").lower() if para.style is not None else ""
            level = re.search(r"heading (\d)", style)
            if style == "title":
                out.append(f"# {para.text.strip()}")
            elif level:
                out.append(f"{'#' * (int(level.group(1)) + 1)} {para.text.strip()}")
            elif "list" in style:
                out.append(f"- {text}")
            else:
                out.append(text)
            out.append("")
        elif tag == "tbl":
            table = docx.table.Table(block, doc)
            rows = [[c.text.strip().replace("\n", " ") for c in row.cells] for row in table.rows]
            if rows:
                out.append("| " + " | ".join(rows[0]) + " |")
                out.append("|" + "---|" * len(rows[0]))
                out += ["| " + " | ".join(r) + " |" for r in rows[1:]]
                out.append("")
    return "\n".join(out).strip() + "\n"


def docx_to_html(src: Path, images_dir: Path) -> str:
    """Word -> HTML keeping headings, lists, tables, bold/italic and pictures (saved in
    images_dir). Used for Word -> PDF when neither Office nor LibreOffice is installed."""
    import mammoth

    counter = [0]

    def save_image(image):
        counter[0] += 1
        name = f"img{counter[0]}.{(image.content_type or 'image/png').split('/')[-1].replace('jpeg', 'jpg')}"
        with image.open() as data:
            (images_dir / name).write_bytes(data.read())
        return {"src": name}

    with open(src, "rb") as f:
        result = mammoth.convert_to_html(f, convert_image=mammoth.images.img_element(save_image))
    return (f"<!doctype html><html><head><meta charset='utf-8'><style>{PAGE_CSS}</style></head>"
            f"<body>{result.value}</body></html>")


def text_as_markdown(src: Path) -> str:
    """Readable Markdown for any supported input (OCR for images and scanned PDFs)."""
    s = src.suffix.lower()
    if s in MARKDOWN | TEXT:
        return src.read_text(encoding="utf-8", errors="replace")
    if s in HTML:
        from .sources import html_to_text

        title, text = html_to_text(src.read_text(encoding="utf-8", errors="replace"))
        return (f"# {title}\n\n" if title else "") + text.replace("\n", "\n\n")
    if s in WORD:
        return docx_to_markdown(src)
    if s in IMAGES:
        from . import ocr

        return ocr.ocr_document(src.name, src.read_bytes())[0]["text"]
    if s in PDF:
        from . import ocr

        pages = ocr.ocr_document(src.name, src.read_bytes()) if ocr.available() else None
        if pages is None:
            import pymupdf

            with pymupdf.open(src) as doc:
                pages = [{"page": i + 1, "text": p.get_text("text", sort=True)} for i, p in enumerate(doc)]
        if len(pages) == 1:
            return pages[0]["text"].strip() + "\n"
        return "\n\n".join(f"## Page {p['page']}\n\n{p['text'].strip()}" for p in pages) + "\n"
    if s in OFFICE:
        with tempfile.TemporaryDirectory() as tmp:
            done = _office_convert(src, "docx", Path(tmp) / "doc.docx")
            if not done:
                raise ConvertError(f"Reading {s} files needs Microsoft Office or LibreOffice (free): "
                                   "https://www.libreoffice.org/download/")
            return docx_to_markdown(done)
    raise ConvertError(f"{src.name}: unsupported file type")


def images_to_pdf(images: list[Path], out: Path) -> Path:
    import pymupdf

    doc = pymupdf.open()
    for path in images:
        with pymupdf.open(path) as img:
            pdf_bytes = img.convert_to_pdf()
        with pymupdf.open("pdf", pdf_bytes) as page_doc:
            doc.insert_pdf(page_doc)
    doc.save(out, deflate=True)
    return out


def pdf_to_images(pdf: Path, out: Path, dpi: int = 150) -> Path:
    """Save every page as PNG: one .png for a single page, otherwise a .zip of them."""
    import pymupdf

    with pymupdf.open(pdf) as doc:
        if len(doc) == 1:
            out = out.with_suffix(".png")
            doc[0].get_pixmap(dpi=dpi).save(out)
            return out
        out = out.with_suffix(".zip")
        with zipfile.ZipFile(out, "w") as zf:
            for i, page in enumerate(doc, 1):
                zf.writestr(f"{pdf.stem}_page{i:03}.png", page.get_pixmap(dpi=dpi).tobytes("png"))
    return out


def merge_pdfs(files: list[Path], out: Path, work_dir: Path) -> Path:
    """Join files into one PDF; non-PDF files are converted first."""
    import pymupdf

    merged = pymupdf.open()
    for f in files:
        pdf = f if f.suffix.lower() in PDF else convert(f, "pdf", work_dir)
        with pymupdf.open(pdf) as doc:
            merged.insert_pdf(doc)
    merged.save(out, deflate=True)
    return out


def searchable_pdf(src: Path, out: Path, *, step=None, should_stop=None, dpi: int = 200) -> Path:
    """Scanned PDF or photo -> PDF with an invisible OCR text layer you can search and copy."""
    import pymupdf

    from . import ocr

    if ocr.available() != "rapidocr":
        raise ConvertError("Searchable PDFs need the built-in OCR: pip install rapidocr onnxruntime")
    if src.suffix.lower() in IMAGES:
        with pymupdf.open(src) as img:
            source = pymupdf.open("pdf", img.convert_to_pdf())
    else:
        source = pymupdf.open(src)
    result = pymupdf.open()
    result.insert_pdf(source)
    for i, page in enumerate(result):
        if should_stop and should_stop():
            raise InterruptedError()
        if step:
            step(i + 1, len(result))
        if len(page.get_text().strip()) >= 25:
            continue  # already has real text
        scale = 72 / dpi
        for box, text in ocr.ocr_boxes(page.get_pixmap(dpi=dpi).tobytes("png")):
            xs, ys = [p[0] * scale for p in box], [p[1] * scale for p in box]
            x0, y0, x1, y1 = min(xs), min(ys), max(xs), max(ys)
            height = max(y1 - y0, 4)
            font = "china-s" if re.search(r"[\u3400-\u9fff]", text) else "helv"
            size = height * 0.85
            width = pymupdf.get_text_length(text, fontname=font, fontsize=size) or 1
            origin = pymupdf.Point(x0, y1 - height * 0.2)
            page.insert_text(origin, text, fontsize=size, fontname=font, render_mode=3,
                             morph=(origin, pymupdf.Matrix((x1 - x0) / width, 1)))
    result.save(out, garbage=3, deflate=True)
    return out


# --- Main entry points ------------------------------------------------------------------

def convert(src: Path, target: str, out_dir: Path, *, notes: list[str] | None = None,
            step=None, should_stop=None) -> Path:
    """Convert one file, keeping its look where possible. Returns the new file in out_dir.

    notes collects remarks for the user (e.g. when a layout could only be approximated);
    step(n, total) and should_stop() report progress and allow stopping during OCR.
    """
    notes = notes if notes is not None else []
    s = src.suffix.lower()
    if s not in INPUTS:
        raise ConvertError(f"{src.name}: unsupported file type")
    if target not in TARGETS:
        raise ConvertError(f"Unknown target format: {target}")
    out = out_dir / f"{src.stem}.{target}"
    if out.resolve() == src.resolve():
        out = out_dir / f"{src.stem} (converted).{target}"

    if target == "pdf":
        if s in PDF:
            return Path(shutil.copy(src, out))
        if s in IMAGES:
            return images_to_pdf([src], out)
        if s in WORD | OFFICE:
            done = _office_convert(src, "pdf", out)
            if done:
                return done
            if s in OFFICE:
                raise ConvertError(f"Converting {s} files needs Microsoft Office or LibreOffice (free): "
                                   "https://www.libreoffice.org/download/")
            with tempfile.TemporaryDirectory() as tmp:
                html_to_pdf(docx_to_html(src, Path(tmp)), out, images_dir=Path(tmp))
            notes.append(f"{src.name}: laid out without Word or LibreOffice, so fonts and page "
                         "breaks may differ. Install LibreOffice (free) for an exact copy.")
            return out
        if s in HTML:
            return html_to_pdf(src.read_text(encoding="utf-8", errors="replace"), out)
        return html_to_pdf(markdown_to_html(text_as_markdown(src), src.stem), out)
    if target == "docx":
        if s in WORD:
            return Path(shutil.copy(src, out))
        if s in OFFICE:
            done = _office_convert(src, "docx", out)
            if not done:
                raise ConvertError(f"Converting {s} files needs Microsoft Office or LibreOffice (free): "
                                   "https://www.libreoffice.org/download/")
            return done
        if s in PDF | IMAGES:
            return pdf_to_docx(src, out, notes=notes, step=step, should_stop=should_stop)
        return markdown_to_docx(text_as_markdown(src), out, src.stem)
    if target in ("txt", "md"):
        text = text_as_markdown(src)
        if target == "txt":
            text = re.sub(r"\*\*(.+?)\*\*", r"\1", text)
        out.write_text(text, encoding="utf-8")
        return out
    if target == "html":
        if s in HTML:
            return Path(shutil.copy(src, out))
        out.write_text(markdown_to_html(text_as_markdown(src), src.stem), encoding="utf-8")
        return out
    # png
    pdf = src if s in PDF else convert(src, "pdf", out_dir, notes=notes)
    return pdf_to_images(pdf, out)


def pdf_to_docx(src: Path, out: Path, *, notes: list[str] | None = None, step=None,
                should_stop=None) -> Path:
    """PDF or image -> Word with the same layout. Scanned pages are first rebuilt with
    their text in place (OCR), so the Word file is editable and looks like the original."""
    from . import ocr

    notes = notes if notes is not None else []
    try:
        from pdf2docx import Converter
    except ImportError:
        Converter = None
    with tempfile.TemporaryDirectory() as tmp:
        pdf = src
        scanned = src.suffix.lower() in IMAGES or not _pdf_has_text(src, every_page=True)
        if scanned and ocr.available():
            from .layout import layout_pdf

            pdf = layout_pdf(src, Path(tmp) / "layout.pdf", step=step, should_stop=should_stop)
        elif src.suffix.lower() in IMAGES:
            pdf = images_to_pdf([src], Path(tmp) / "image.pdf")
        if Converter is None:
            notes.append("pdf2docx is not installed, so only the text was kept.")
            return markdown_to_docx(text_as_markdown(src), out, src.stem)
        conv = Converter(str(pdf))
        try:
            conv.convert(str(out))
        finally:
            conv.close()
    return out


def _pdf_has_text(path: Path, every_page: bool = False) -> bool:
    """Whether the PDF has selectable text (on every page that has an image, if every_page)."""
    import pymupdf

    with pymupdf.open(path) as doc:
        if every_page:
            return all(len(p.get_text().strip()) >= 25 or not p.get_images() for p in doc)
        return sum(len(p.get_text().strip()) for p in doc) > 50 * max(len(doc), 1) * 0.2


def export_markdown(text: str, fmt: str, title: str = "",
                    images_dir: Path | None = None) -> tuple[bytes, str]:
    """Markdown -> (file bytes, media type) for downloads. Pictures (![..](name)) are read
    from images_dir for PDF and Word."""
    if fmt == "md":
        return text.encode("utf-8"), "text/markdown"
    if fmt == "txt":
        plain = re.sub(r"\*\*(.+?)\*\*", r"\1", re.sub(r"(?m)^#{1,6}\s+", "", text))
        return plain.encode("utf-8"), "text/plain"
    if fmt == "html":
        return markdown_to_html(text, title).encode("utf-8"), "text/html"
    with tempfile.TemporaryDirectory() as tmp:
        out = Path(tmp) / f"export.{fmt}"
        if fmt == "pdf":
            html_to_pdf(markdown_to_html(text, title), out, header=title, images_dir=images_dir)
            return out.read_bytes(), "application/pdf"
        if fmt == "docx":
            markdown_to_docx(text, out, title, images_dir)
            return out.read_bytes(), ("application/vnd.openxmlformats-officedocument."
                                      "wordprocessingml.document")
    raise ConvertError(f"Unknown export format: {fmt}")


def main(argv: list[str] | None = None) -> int:
    """python -m sgpt.convert report.docx --to pdf"""
    import argparse
    import sys

    p = argparse.ArgumentParser(prog="sgpt.convert", description="Convert documents.")
    p.add_argument("files", nargs="+", help="files to convert")
    p.add_argument("--to", required=True, choices=[*TARGETS, "merged-pdf", "searchable-pdf"],
                   help="target format; merged-pdf joins all files into one PDF")
    p.add_argument("-o", "--out-dir", default=".", help="where to save (default: current folder)")
    args = p.parse_args(argv)
    out_dir = Path(args.out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)
    try:
        if args.to == "merged-pdf":
            print(merge_pdfs([Path(f) for f in args.files], out_dir / "merged.pdf", out_dir))
            return 0
        for f in args.files:
            src = Path(f)
            if args.to == "searchable-pdf":
                print(searchable_pdf(src, out_dir / f"{src.stem} (searchable).pdf"))
            else:
                print(convert(src, args.to, out_dir))
    except (ConvertError, OSError) as exc:
        print(f"Error: {exc}", file=sys.stderr)
        return 1
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
